# analise_remocao.py - Versão 14.0 - Análise de Remoção de Magistrados com Lógica de "Congelamento"

import pandas as pd
from docx import Document
import os
import networkx as nx
from pyvis.network import Network
from openpyxl.styles import Alignment

# --- CONFIGURAÇÕES ---
VAGAS_INICIAIS = [
    "22A VARA DA SJ MINAS GERAIS",
    "VARA ÚNICA DA SSJ MURIAÉ"
]

def ler_arquivo_word(caminho_arquivo):
    if not os.path.exists(caminho_arquivo): return None
    try:
        doc = Document(caminho_arquivo)
        if len(doc.tables) == 0: return None
        tabela = doc.tables[0]
        dados = []
        chaves = []
        for i, linha in enumerate(tabela.rows):
            textos = [c.text.strip() for c in linha.cells]
            if i == 0:
                chaves = textos
                continue
            dados.append(dict(zip(chaves, textos)))
        return pd.DataFrame(dados)
    except Exception as e:
        print(f"Erro: {e}")
        return None

def normalizar_nome(texto):
    if not isinstance(texto, str): return ""
    texto = texto.upper().strip()
    if " - " in texto:
        texto = texto.split(" - ")[0]
    return texto.strip()

def encurtar_nome(texto):
    """Deixa os nomes das varas menores para caber na bolinha do gráfico."""
    if not texto: return "Origem Desconhecida"
    texto = normalizar_nome(texto)
    texto = texto.replace("DA SJ ", "").replace("DA SSJ ", "").replace("SUBSEÇÃO JUDICIÁRIA DE ", "")
    texto = texto.replace("SEÇÃO JUDICIÁRIA DE ", "").replace("SEÇÃO JUDICIÁRIA DO ", "")
    texto = texto.replace("VARA ÚNICA", "VARA UNICA").replace("RELATORIA DA ", "REL. ")
    if len(texto) > 35:
        return texto[:32] + "..."
    return texto

def gerar_grafo_visual(df_resultado, pasta_destino):
    """Gera um arquivo HTML interativo com a visualização da cadeia."""
    print("\n--- GERANDO GRÁFICO VISUAL (GRAFO) ---")
    # Fundo escuro (#222222)
    net = Network(height='750px', width='100%', bgcolor='#222222', font_color='white', directed=True)
    net.barnes_hut(gravity=-3000, central_gravity=0.3, spring_length=200)

    for index, row in df_resultado.iterrows():
        partes_nome = row['Nome'].split(" ")
        juiz_curto = partes_nome[0] + " " + partes_nome[-1] 
        origem = encurtar_nome(row['Origem'])
        destino = encurtar_nome(row['Destino'])
        if not origem: origem = "Externo"
            
        net.add_node(origem, label=origem, color='#ff6b6b', title="Origem/Vaga Aberta") 
        net.add_node(destino, label=destino, color='#51cf66', title="Destino Ocupado")
        
        # --- CORREÇÃO VISUAL ---
        # Adicionamos 'background': '#222222' (mesma cor do fundo)
        # Isso faz parecer que a linha se rompe onde tem texto.
        net.add_edge(
            origem, 
            destino, 
            title=f"Magistrado: {row['Nome']}", 
            label=juiz_curto, 
            color='white',
            font={
                'size': 9, 
                'align': 'middle', 
                'color': 'white', 
                'background': '#222222', # O Pulo do Gato: Fundo da etiqueta igual ao fundo da tela
                'strokeWidth': 0
            } 
        )

    caminho_html = os.path.join(pasta_destino, "Grafo_Remocao.html")
    try:
        net.save_graph(caminho_html)
        print(f"📊 Gráfico interativo gerado com sucesso: {caminho_html}")
    except Exception as e:
        print(f"Erro ao salvar gráfico HTML: {e}")
    return caminho_html

def processar_remocao_com_repescagem(df):
    print("\n--- INICIANDO ALGORITMO DE REPESCAGEM (RESTART ON MOVE) ---")
    vagas_abertas = set([normalizar_nome(v) for v in VAGAS_INICIAIS])
    remocoes_confirmadas = {} 
    log_movimentacoes = []
    ciclo = 0
    houve_movimentacao = True
    
    while houve_movimentacao:
        houve_movimentacao = False
        ciclo += 1
        for index, juiz in df.iterrows():
            matricula = juiz['Matrícula']
            if matricula in remocoes_confirmadas: continue
            
            nome = juiz['Nome']
            lotacao_atual = juiz['Lotação Atual']
            lotacao_atual_norm = normalizar_nome(lotacao_atual)
            opcoes = juiz['Lista_Opcoes']
            
            match_encontrado = False
            vaga_escolhida = ""
            opcao_numero = 0
            opcao_final_texto = ""
            
            for i, opcao_bruta in enumerate(opcoes):
                opcao_norm = normalizar_nome(opcao_bruta)
                for vaga_aberta in vagas_abertas:
                    if vaga_aberta in opcao_norm:
                        match_encontrado = True
                        vaga_escolhida = vaga_aberta
                        opcao_numero = i + 1
                        opcao_final_texto = opcao_norm
                        break
                if match_encontrado: break
            
            if match_encontrado:
                remocoes_confirmadas[matricula] = {
                    'Matrícula': matricula,
                    'Nome': nome,
                    'Origem': lotacao_atual,
                    'Destino': opcao_final_texto,
                    'Opção Nº': opcao_numero,
                    'Ciclo': ciclo
                }
                log_movimentacoes.append(f"✅ CICLO {ciclo}: {nome} assumiu {opcao_final_texto[:40]}...")
                vagas_abertas.remove(vaga_escolhida)
                if lotacao_atual_norm:
                    vagas_abertas.add(lotacao_atual_norm)
                    log_movimentacoes.append(f"   -> Abriu vaga: {lotacao_atual_norm[:40]}...")
                houve_movimentacao = True
                break 
    
    return pd.DataFrame(list(remocoes_confirmadas.values())), log_movimentacoes, vagas_abertas

# --- EXECUÇÃO PRINCIPAL ---
nome_arquivo = "Quadro de Magistrados Inscritos para Remoção - 21.11 - 09h.docx"
pasta_atual = os.getcwd()
caminho_completo = os.path.join(pasta_atual, nome_arquivo)

print(f"Lendo arquivo em: {pasta_atual}")
df_bruto = ler_arquivo_word(caminho_completo)

if df_bruto is not None:
    df_bruto['Data de Exercício'] = pd.to_datetime(df_bruto['Data de Exercício'], format='%d/%m/%Y', errors='coerce')
    df_ordenado = df_bruto.sort_values(by=['Data de Exercício', 'Matrícula'], ascending=[True, True])
    
    opcoes_limpas = []
    colunas_opcoes = [col for col in df_bruto.columns if 'Opção' in col]
    for _, row in df_ordenado.iterrows():
        lista = [row[col] for col in colunas_opcoes if row[col] and str(row[col]).strip() != ""]
        opcoes_limpas.append(lista)
    df_ordenado['Lista_Opcoes'] = opcoes_limpas
    
    # Processamento
    df_resultado, log, sobras = processar_remocao_com_repescagem(df_ordenado)
    
    print("\n" + "="*50)
    print("       RESULTADO DA PROJEÇÃO")
    print("="*50)
    
    if not df_resultado.empty:
        for linha in log:
            print(linha)
            
        print("\n--- TABELA FINAL (RESUMO) ---")
        pd.set_option('display.max_colwidth', None)
        print(df_resultado[['Nome', 'Destino', 'Opção Nº']])

        arquivo_saida = os.path.join(pasta_atual, "Resultado_Final_Repescagem.xlsx")
        df_sobras = pd.DataFrame(list(sobras), columns=["Vagas que Sobraram"])
        
        try:
            with pd.ExcelWriter(arquivo_saida, engine='openpyxl') as writer:
                df_resultado.to_excel(writer, index=False, sheet_name='Resultado')
                df_sobras.to_excel(writer, index=False, sheet_name='Resultado', startcol=8)
                
                worksheet = writer.sheets['Resultado']
                estilo_esquerda = Alignment(horizontal='left')
                worksheet['I1'].alignment = estilo_esquerda
                for i in range(len(df_sobras)):
                    worksheet[f'I{i+2}'].alignment = estilo_esquerda
                
            print(f"\n📁 Excel gerado: {arquivo_saida}")
            
        except Exception as e:
            print(f"Erro ao salvar Excel: {e}")
        
        gerar_grafo_visual(df_resultado, pasta_atual)
            
    else:
        print("⚠️ Nenhuma remoção foi realizada.")

    print("\n--- VAGAS QUE SOBRARAM ---")
    for v in sobras:
        print(f"- {v}")

else:
    print("❌ Falha na leitura inicial.")